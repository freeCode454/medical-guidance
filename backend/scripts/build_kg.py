"""从知识库源数据 Excel 构建医学分诊知识图谱（Neo4j）。

图模型：

    节点：Department(科室) / Disease(疾病) / Symptom(症状) / Drug(药物) / ICD(编码) / Condition(禁忌条件)
    关系：
      (Disease)-[:BELONGS_TO]->(Department)       疾病归属科室
      (Symptom)-[:INDICATES]->(Disease)           症状提示疾病
      (Symptom)-[:SUGGESTS_DEPT]->(Department)    症状提示科室（分诊规则）
      (Disease)-[:CODED_AS]->(ICD)                疾病对应 ICD-10 编码
      (Drug)-[:CONTRAINDICATED_FOR]->(Condition)  药物禁忌（带严重度）
      (Symptom)-[:SYNONYM_OF]->(Symptom)          口语同义词 → 标准症状

数据全部来自 kb/医学知识库源数据.xlsx，不重新标注。MERGE 幂等，可重复运行。

用法（backend 目录）：
    python scripts/build_kg.py
环境变量（可选）：NEO4J_URI / NEO4J_USER / NEO4J_PASSWORD
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

from neo4j import GraphDatabase
from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parent.parent
EXCEL = BACKEND_DIR.parent / "kb" / "医学知识库源数据.xlsx"

NEO4J_URI = os.getenv("NEO4J_URI", "bolt://127.0.0.1:7687")
NEO4J_USER = os.getenv("NEO4J_USER", "neo4j")
NEO4J_PASSWORD = os.getenv("NEO4J_PASSWORD", "12345678")


def _rows(sheet: str) -> list[list]:
    ws = load_workbook(EXCEL)[sheet]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            out.append([str(c) if c is not None else "" for c in row])
    return out


def _split(s: str) -> list[str]:
    return [x.strip() for x in s.split("|") if x.strip()]


def build() -> dict[str, list[dict]]:
    """抽取节点与关系（返回按类型分组的批量数据）。"""
    depts = _rows("科室字典")
    diseases = _rows("疾病特征库")
    symptoms = _rows("症状本体库")
    triage = _rows("分诊规则")
    icd = _rows("ICD-10映射")
    drugs = _rows("禁忌症与用药风险")

    # ---- 节点 ----
    dept_nodes = [{"code": d[0], "name": d[1], "aliases": d[2]} for d in depts]
    # 疾病特征库列顺序：疾病名称, 科室编码, 症状, 鉴别要点, 高危因素
    disease_nodes = [{"name": d[0], "dept_code": d[1]} for d in diseases]
    symptom_nodes = [{"name": s[0], "body_part": s[2]} for s in symptoms]
    icd_nodes = [{"code": d[1], "chapter": d[2], "disease": d[0]} for d in icd]
    condition_names = sorted({d[1] for d in drugs if d[1]})
    condition_nodes = [{"name": n} for n in condition_names]

    # ---- 关系 ----
    belongs_to = [{"disease": d[0], "dept_code": d[1]} for d in diseases]
    # 症状字段按「核心症状在前」排列：首个记 1.0，其余记 0.6（伴随症状）。
    # 用于区分「胸痛之于心绞痛」（核心）与「胸痛之于肺癌」（伴随）。
    indicates = [
        {"symptom": sym, "disease": d[0], "weight": 1.0 if idx == 0 else 0.6}
        for d in diseases
        for idx, sym in enumerate(_split(d[2]))
    ]
    suggests_dept = [
        {"symptom": sym, "dept_code": t[1], "priority": int(float(t[5] or 8))}
        for t in triage
        for sym in _split(t[0])
    ]
    coded_as = [{"disease": d[0], "icd_code": d[1]} for d in icd]
    contraindicated = [
        {"drug": d[0], "condition": d[1], "severity": d[2]} for d in drugs if d[1]
    ]
    synonyms = [
        {"synonym": syn, "standard": s[0]}
        for s in symptoms
        for syn in _split(s[1])
    ]

    return {
        "dept": dept_nodes,
        "disease": disease_nodes,
        "symptom": symptom_nodes,
        "icd": icd_nodes,
        "condition": condition_nodes,
        "belongs_to": belongs_to,
        "indicates": indicates,
        "suggests_dept": suggests_dept,
        "coded_as": coded_as,
        "contraindicated": contraindicated,
        "synonym": synonyms,
    }


# 同一 Neo4j 实例里可能还有别的图谱（如设备维修），全局唯一约束会互相冲突，
# 因此本项目不用唯一约束，靠 MERGE 保证幂等；节点统一打 :Mg 标签隔离。
SCHEMA = [
    "DROP CONSTRAINT dept_code IF EXISTS",
    "DROP CONSTRAINT disease_name IF EXISTS",
    "DROP CONSTRAINT symptom_name IF EXISTS",
    "DROP CONSTRAINT drug_name IF EXISTS",
    "DROP CONSTRAINT icd_code IF EXISTS",
    "DROP CONSTRAINT cond_name IF EXISTS",
]

BATCH = 500

# 本项目所有节点统一打 :Mg 标签，避免与同一 Neo4j 实例里的其他图谱混淆
MG = ":Mg"


def _chunks(rows: list, size: int = BATCH):
    for i in range(0, len(rows), size):
        yield rows[i : i + size]


def write(session, data: dict[str, list[dict]]) -> None:
    for cql in SCHEMA:
        session.run(cql)

    # 清理：首次运行遗留的医学节点（无 :Mg）+ 本项目已有数据
    session.run(
        "MATCH (n) WHERE (n:Department OR n:Disease OR n:ICD OR n:Condition) "
        "AND NOT n:Mg DETACH DELETE n"
    )
    session.run("MATCH (n:Mg) DETACH DELETE n")

    # 节点
    for batch in _chunks(data["dept"]):
        session.run(
            f"UNWIND $rows AS r MERGE (n:Department{MG} {{code: r.code}}) "
            "SET n.name = r.name, n.aliases = r.aliases",
            rows=batch,
        )
    for batch in _chunks(data["disease"]):
        session.run(
            f"UNWIND $rows AS r MERGE (n:Disease{MG} {{name: r.name}}) SET n.dept_code = r.dept_code",
            rows=batch,
        )
    for batch in _chunks(data["symptom"]):
        session.run(
            f"UNWIND $rows AS r MERGE (n:Symptom{MG} {{name: r.name}}) SET n.body_part = r.body_part",
            rows=batch,
        )
    for batch in _chunks(data["icd"]):
        session.run(
            f"UNWIND $rows AS r MERGE (n:ICD{MG} {{code: r.code}}) SET n.chapter = r.chapter",
            rows=batch,
        )
    for batch in _chunks(data["condition"]):
        session.run(f"UNWIND $rows AS r MERGE (n:Condition{MG} {{name: r.name}})", rows=batch)
    drug_rows = [{"name": d["drug"]} for d in data["contraindicated"]]
    for batch in _chunks(drug_rows):
        session.run(f"UNWIND $rows AS r MERGE (n:Drug{MG} {{name: r.name}})", rows=batch)

    # 关系
    for batch in _chunks(data["belongs_to"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (d:Disease:Mg {name: r.disease}), (dept:Department:Mg {code: r.dept_code}) "
            "MERGE (d)-[:BELONGS_TO]->(dept)",
            rows=batch,
        )
    for batch in _chunks(data["indicates"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (s:Symptom:Mg {name: r.symptom}), (d:Disease:Mg {name: r.disease}) "
            "MERGE (s)-[rel:INDICATES]->(d) SET rel.weight = r.weight",
            rows=batch,
        )
    for batch in _chunks(data["suggests_dept"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (s:Symptom:Mg {name: r.symptom}), (dept:Department:Mg {code: r.dept_code}) "
            "MERGE (s)-[rel:SUGGESTS_DEPT]->(dept) SET rel.priority = r.priority",
            rows=batch,
        )
    for batch in _chunks(data["coded_as"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (d:Disease:Mg {name: r.disease}), (i:ICD:Mg {code: r.icd_code}) "
            "MERGE (d)-[:CODED_AS]->(i)",
            rows=batch,
        )
    for batch in _chunks(data["contraindicated"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (m:Drug:Mg {name: r.drug}), (c:Condition:Mg {name: r.condition}) "
            "MERGE (m)-[rel:CONTRAINDICATED_FOR]->(c) SET rel.severity = r.severity",
            rows=batch,
        )
    for batch in _chunks(data["synonym"]):
        session.run(
            "UNWIND $rows AS r "
            "MATCH (std:Symptom:Mg {name: r.standard}) "
            "MERGE (syn:Symptom:Mg {name: r.synonym}) "
            "MERGE (syn)-[:SYNONYM_OF]->(std)",
            rows=batch,
        )


def main() -> None:
    if not EXCEL.exists():
        print(f"未找到 {EXCEL}，请先运行 gen_kb_excel.py")
        return
    data = build()
    driver = GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))
    try:
        with driver.session(database="neo4j") as session:
            write(session, data)
            counts = session.run(
                "MATCH (n) RETURN labels(n)[0] AS label, count(*) AS c ORDER BY c DESC"
            ).data()
            rels = session.run(
                "MATCH ()-[r]->() RETURN type(r) AS type, count(*) AS c ORDER BY c DESC"
            ).data()
        print("节点：")
        for c in counts:
            print(f"  {c['label']}: {c['c']}")
        print("关系：")
        for r in rels:
            print(f"  {r['type']}: {r['c']}")
    finally:
        driver.close()


if __name__ == "__main__":
    main()
